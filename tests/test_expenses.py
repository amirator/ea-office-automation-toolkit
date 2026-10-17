from openpyxl import load_workbook

from conftest import SAMPLE
from toolkit.expenses import generate, load, summarise


def data():
    return load(SAMPLE / "expenses.csv", SAMPLE / "trips.csv", SAMPLE / "fx_rates.csv", SAMPLE / "policy_limits.csv")


def test_currency_conversion_to_sar():
    lines = data().lines
    hotel_london = next(ln for ln in lines if ln.description == "Hotel in the City of London")
    assert hotel_london.sar == round(1040 * 4.98, 2)
    assert hotel_london.sar_per_unit == round(1040 * 4.98 / 4, 2)


def test_policy_and_receipt_flags():
    s = summarise(data())
    over = {ln.description for ln in s["over_limit"]}
    assert "Hotel in the City of London" in over          # SAR 1,295/night > 1,100
    assert "Hotel in Business Bay" not in over            # SAR 970/night
    missing = {ln.description for ln in s["missing_receipt"]}
    assert "Ride-hailing between bank meetings" in missing  # SAR 180, no receipt
    assert all(ln.sar > 100 for ln in s["missing_receipt"])


def test_trip_totals_add_up():
    s = summarise(data())
    assert round(sum(t["actual_sar"] for t in s["trips"]), 2) == s["total_sar"]
    assert round(sum(s["by_category"].values()), 2) == s["total_sar"]


def test_workbook_structure_and_live_formulas(tmp_path):
    res = generate(SAMPLE / "expenses.csv", SAMPLE / "trips.csv", SAMPLE / "fx_rates.csv",
                   SAMPLE / "policy_limits.csv", tmp_path)
    wb = load_workbook(res["workbook"])
    assert wb.sheetnames == ["Summary", "Expenses", "Trips", "Policy", "FX Rates"]
    assert str(wb["Expenses"]["K5"].value).startswith("=ROUND(")
    assert "SUMIFS" in str(wb["Trips"]["I5"].value)
    assert "SAMPLE" in wb["Summary"]["A2"].value
    assert len(wb["Summary"]._charts) == 2
    assert "SAR" in res["summary"].read_text()
