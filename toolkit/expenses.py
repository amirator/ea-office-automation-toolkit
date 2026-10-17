"""Expense / travel tracker: builds a formula-driven Excel workbook and a short text summary.

The workbook keeps live Excel formulas (FX conversion, policy look-ups, flags,
per-trip SUMIFS) so finance can keep using it after the script has run. The
same numbers are also computed in Python for the text summary and the tests.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from toolkit.common import GREY, LIGHT, NAVY, RAG_COLOURS, SAMPLE_BANNER, TEAL, ensure_dir, read_csv

RECEIPT_THRESHOLD_SAR = 100  # receipts required above this amount


@dataclass
class Line:
    date: date
    trip_id: str
    category: str
    description: str
    amount: float
    currency: str
    units: int
    payment_method: str
    receipt: str
    fx: float
    limit: float
    per: str

    @property
    def sar(self) -> float:
        return round(self.amount * self.fx, 2)

    @property
    def sar_per_unit(self) -> float:
        return round(self.sar / max(self.units, 1), 2)

    @property
    def flags(self) -> list[str]:
        f = []
        if self.sar_per_unit > self.limit:
            f.append("Over limit")
        if self.receipt.upper() != "Y" and self.sar > RECEIPT_THRESHOLD_SAR:
            f.append("Missing receipt")
        return f


@dataclass
class ExpenseData:
    lines: list[Line]
    trips: list[dict]
    rates: dict[str, float]
    policy: list[dict]


def load(expenses_csv, trips_csv, rates_csv, policy_csv) -> ExpenseData:
    rates = {r["currency"].upper(): float(r["sar_per_unit"]) for r in read_csv(rates_csv)}
    policy = read_csv(policy_csv)
    limits = {p["category"]: (float(p["limit_sar"]), p["per"]) for p in policy}
    trips = read_csv(trips_csv)
    trip_ids = {t["trip_id"] for t in trips}
    lines = []
    for n, r in enumerate(read_csv(expenses_csv), 2):
        cur = r["currency"].upper()
        if cur not in rates:
            raise ValueError(f"Row {n}: no FX rate for currency {cur}")
        if r["category"] not in limits:
            raise ValueError(f"Row {n}: category '{r['category']}' is not in the policy table")
        if r["trip_id"] not in trip_ids:
            raise ValueError(f"Row {n}: unknown trip {r['trip_id']}")
        lim, per = limits[r["category"]]
        lines.append(Line(date.fromisoformat(r["date"]), r["trip_id"], r["category"], r["description"],
                          float(r["amount"]), cur, int(r.get("units") or 1), r.get("payment_method", ""),
                          r.get("receipt", "N"), rates[cur], lim, per))
    return ExpenseData(lines, trips, rates, policy)


def summarise(data: ExpenseData) -> dict:
    by_cat, by_trip = defaultdict(float), defaultdict(float)
    for ln in data.lines:
        by_cat[ln.category] += ln.sar
        by_trip[ln.trip_id] += ln.sar
    total = round(sum(by_cat.values()), 2)
    flagged = [ln for ln in data.lines if ln.flags]
    trips = []
    for t in data.trips:
        actual = round(by_trip.get(t["trip_id"], 0.0), 2)
        budget = float(t["budget_sar"])
        trips.append({**t, "actual_sar": actual, "budget_sar": budget,
                      "variance_sar": round(budget - actual, 2), "pct_used": actual / budget if budget else 0})
    return {
        "total_sar": total,
        "by_category": dict(sorted(((k, round(v, 2)) for k, v in by_cat.items()), key=lambda kv: -kv[1])),
        "trips": trips,
        "flagged": flagged,
        "line_count": len(data.lines),
        "over_limit": [ln for ln in flagged if "Over limit" in ln.flags],
        "missing_receipt": [ln for ln in flagged if "Missing receipt" in ln.flags],
        "currencies": sorted({ln.currency for ln in data.lines}),
        "period": (min(ln.date for ln in data.lines), max(ln.date for ln in data.lines)),
    }


def text_summary(s: dict) -> str:
    top_cat, top_val = next(iter(s["by_category"].items()))
    over = [t for t in s["trips"] if t["actual_sar"] > t["budget_sar"]]
    p0, p1 = s["period"]
    parts = [
        f"{len(s['trips'])} trips between {p0:%d %b} and {p1:%d %b %Y} cost SAR {s['total_sar']:,.0f} in total "
        f"across {sum(1 for _ in s['by_category'])} categories and {len(s['currencies'])} currencies "
        f"({', '.join(s['currencies'])}).",
        f"{top_cat} is the largest category at SAR {top_val:,.0f} ({top_val / s['total_sar']:.0%} of spend).",
    ]
    if over:
        parts.append("Over budget: " + "; ".join(
            f"{t['trip_id']} {t['destination']} by SAR {-t['variance_sar']:,.0f} ({t['pct_used']:.0%} of budget)"
            for t in over) + ".")
    else:
        parts.append("All trips are within budget.")
    parts.append(f"{len(s['flagged'])} of {s['line_count']} lines need review: "
                 f"{len(s['over_limit'])} above the per-unit policy limit and {len(s['missing_receipt'])} "
                 f"missing a receipt above SAR {RECEIPT_THRESHOLD_SAR}.")
    if s["over_limit"]:
        worst = max(s["over_limit"], key=lambda ln: ln.sar_per_unit / ln.limit)
        parts.append(f"Largest exception: {worst.description} ({worst.trip_id}) at SAR {worst.sar_per_unit:,.0f} "
                     f"per {worst.per} against a SAR {worst.limit:,.0f} limit.")
    return " ".join(parts)


# ----------------------------------------------------------------------------- workbook
THIN = Side(style="thin", color="D9DDE3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
HEAD_FONT = Font(bold=True, color="FFFFFF")
SAR_FMT = '#,##0.00 "SAR"'
SAR0_FMT = '#,##0 "SAR"'


def _header(ws, row, headers, widths=None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[row].height = 30


def _banner(ws, text, cols):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=15, color=NAVY)
    ws["A2"] = SAMPLE_BANNER
    ws["A2"].font = Font(bold=True, size=9, color=RAG_COLOURS["Red"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cols)


def build_workbook(data: ExpenseData, path: str | Path, sample: bool = True) -> Path:
    s = summarise(data)
    wb = Workbook()
    summ = wb.active
    summ.title = "Summary"
    exp = wb.create_sheet("Expenses")
    trips = wb.create_sheet("Trips")
    pol = wb.create_sheet("Policy")
    fx = wb.create_sheet("FX Rates")

    # --- FX Rates
    _banner(fx, "FX rates to SAR", 3)
    fx["A3"] = "Fixed illustrative rates for the demo, not live market rates. SAR is pegged at 3.75 per USD."
    fx["A3"].font = Font(italic=True, color=GREY)
    _header(fx, 4, ["Currency", "SAR per unit"], [14, 16])
    for i, (cur, rate) in enumerate(data.rates.items(), 5):
        fx.cell(i, 1, cur).border = BORDER
        c = fx.cell(i, 2, rate)
        c.number_format, c.border = "0.0000", BORDER
    fx_last = 4 + len(data.rates)

    # --- Policy
    _banner(pol, "Travel and expense policy limits", 4)
    pol["A3"] = f"Receipts required for any line above SAR {RECEIPT_THRESHOLD_SAR}."
    pol["A3"].font = Font(italic=True, color=GREY)
    _header(pol, 4, ["Category", "Limit (SAR)", "Per", "Notes"], [26, 14, 10, 60])
    for i, p in enumerate(data.policy, 5):
        for c, v in enumerate([p["category"], float(p["limit_sar"]), p["per"], p.get("notes", "")], 1):
            cell = pol.cell(i, c, v)
            cell.border = BORDER
        pol.cell(i, 2).number_format = SAR0_FMT
    pol_last = 4 + len(data.policy)

    # --- Expenses (formula driven)
    _banner(exp, "Expense lines", 15)
    hdr = ["Date", "Trip ID", "Category", "Description", "Amount", "Currency", "Units", "Payment",
           "Receipt (Y/N)", "FX to SAR", "Amount (SAR)", "SAR per unit", "Policy limit (SAR)", "Limit per", "Flag"]
    _header(exp, 4, hdr, [11, 14, 22, 46, 11, 9, 7, 14, 9, 9, 14, 13, 12, 9, 24])
    first = 5
    for i, ln in enumerate(data.lines, first):
        vals = [ln.date, ln.trip_id, ln.category, ln.description, ln.amount, ln.currency, ln.units,
                ln.payment_method, ln.receipt,
                f"=VLOOKUP(F{i},'FX Rates'!$A$5:$B${fx_last},2,FALSE)",
                f"=ROUND(E{i}*J{i},2)",
                f"=ROUND(K{i}/MAX(G{i},1),2)",
                f"=VLOOKUP(C{i},Policy!$A$5:$C${pol_last},2,FALSE)",
                f"=VLOOKUP(C{i},Policy!$A$5:$C${pol_last},3,FALSE)",
                f'=IF(AND(L{i}>M{i},I{i}<>"Y",K{i}>{RECEIPT_THRESHOLD_SAR}),"Over limit + no receipt",'
                f'IF(L{i}>M{i},"Over limit",IF(AND(I{i}<>"Y",K{i}>{RECEIPT_THRESHOLD_SAR}),"Missing receipt","OK")))']
        for c, v in enumerate(vals, 1):
            cell = exp.cell(i, c, v)
            cell.border = BORDER
        exp.cell(i, 1).number_format = "dd-mmm-yy"
        for c in (6, 7, 9, 14):
            exp.cell(i, c).alignment = Alignment(horizontal="center")
        exp.cell(i, 5).number_format = "#,##0.00"
        exp.cell(i, 10).number_format = "0.0000"
        for c in (11, 12, 13):
            exp.cell(i, c).number_format = "#,##0.00"
    last = first + len(data.lines) - 1
    exp.freeze_panes = "C5"
    exp.auto_filter.ref = f"A4:O{last}"
    red = PatternFill("solid", fgColor="FBE3E3")
    green = PatternFill("solid", fgColor="E3F4EA")
    exp.conditional_formatting.add(f"O{first}:O{last}", FormulaRule(formula=[f'$O{first}<>"OK"'], fill=red,
                                                                     font=Font(bold=True, color="9B1C1C")))
    exp.conditional_formatting.add(f"O{first}:O{last}", FormulaRule(formula=[f'$O{first}="OK"'], fill=green))
    tot = last + 1
    exp.cell(tot, 10, "Total").font = Font(bold=True)
    exp.cell(tot, 11, f"=SUM(K{first}:K{last})").font = Font(bold=True)
    exp.cell(tot, 11).number_format = SAR_FMT
    dv_cat = DataValidation(type="list", formula1=f"=Policy!$A$5:$A${pol_last}", allow_blank=False)
    dv_cur = DataValidation(type="list", formula1=f"='FX Rates'!$A$5:$A${fx_last}", allow_blank=False)
    dv_rec = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    for dv, col in ((dv_cat, "C"), (dv_cur, "F"), (dv_rec, "I")):
        exp.add_data_validation(dv)
        dv.add(f"{col}{first}:{col}{last + 50}")

    # --- Trips
    _banner(trips, "Per-trip totals", 12)
    _header(trips, 4, ["Trip ID", "Traveller", "Purpose", "Destination", "Start", "End", "Nights",
                       "Budget (SAR)", "Actual (SAR)", "Variance (SAR)", "% of budget", "Lines flagged"],
            [14, 18, 44, 12, 11, 11, 8, 14, 14, 15, 11, 12])
    rng = lambda col: f"Expenses!${col}${first}:${col}${last}"
    for i, t in enumerate(data.trips, 5):
        vals = [t["trip_id"], t["traveller"], t["purpose"], t["destination"],
                date.fromisoformat(t["start_date"]), date.fromisoformat(t["end_date"]), f"=F{i}-E{i}",
                float(t["budget_sar"]), f"=SUMIFS({rng('K')},{rng('B')},A{i})", f"=H{i}-I{i}",
                f"=IF(H{i}=0,0,I{i}/H{i})", f'=COUNTIFS({rng("B")},A{i},{rng("O")},"<>OK")']
        for c, v in enumerate(vals, 1):
            trips.cell(i, c, v).border = BORDER
        for c in (5, 6):
            trips.cell(i, c).number_format = "dd-mmm-yy"
        for c in (8, 9, 10):
            trips.cell(i, c).number_format = "#,##0"
        trips.cell(i, 11).number_format = "0%"
    t_last = 4 + len(data.trips)
    trips.conditional_formatting.add(f"J5:J{t_last}", FormulaRule(formula=["$J5<0"], fill=red,
                                                                  font=Font(bold=True, color="9B1C1C")))
    trips.freeze_panes = "B5"

    # --- Summary dashboard
    _banner(summ, "Executive travel and expense summary", 10)
    for col, w in zip("ABCDEFGHIJ", [26, 16, 14, 4, 16, 16, 16, 14, 14, 4]):
        summ.column_dimensions[col].width = w
    tiles = [("Total spend", f"=Expenses!K{tot}", SAR0_FMT),
             ("Trips", f"=COUNTA(Trips!A5:A{t_last})", "0"),
             ("Expense lines", f"=COUNTA(Expenses!A{first}:A{last})", "0"),
             ("Lines flagged", f'=COUNTIF(Expenses!O{first}:O{last},"<>OK")', "0")]
    for n, (label, formula, fmt) in enumerate(tiles):
        col, span = [(1, 1), (2, 3), (5, 6), (7, 8)][n]
        if span > col:
            summ.merge_cells(start_row=4, start_column=col, end_row=4, end_column=span)
            summ.merge_cells(start_row=5, start_column=col, end_row=5, end_column=span)
        lc = summ.cell(4, col, label)
        lc.font, lc.fill = Font(bold=True, color="FFFFFF", size=9), PatternFill("solid", fgColor=TEAL)
        vc = summ.cell(5, col, formula)
        vc.font, vc.number_format = Font(bold=True, size=16, color=NAVY), fmt
        vc.fill = PatternFill("solid", fgColor=LIGHT)
        vc.alignment = Alignment(horizontal="center", vertical="center")
        lc.alignment = Alignment(horizontal="center")
        if n == 3:
            lc.fill = PatternFill("solid", fgColor=RAG_COLOURS["Red"])
    summ.row_dimensions[5].height = 30

    summ["A7"] = "Summary"
    summ["A7"].font = Font(bold=True, size=12, color=NAVY)
    summ["A8"] = text_summary(s)
    summ["A8"].alignment = Alignment(wrap_text=True, vertical="top")
    summ.merge_cells("A8:I10")
    for r in range(8, 11):
        summ.row_dimensions[r].height = 22

    # Spend by category
    _header(summ, 13, ["Category", "Spend (SAR)", "Share"])
    cats = [p["category"] for p in data.policy]
    for i, cat in enumerate(cats, 14):
        summ.cell(i, 1, cat).border = BORDER
        c = summ.cell(i, 2, f"=SUMIFS({rng('K')},{rng('C')},A{i})")
        c.number_format, c.border = "#,##0", BORDER
        c = summ.cell(i, 3, f"=IF(Expenses!$K${tot}=0,0,B{i}/Expenses!$K${tot})")
        c.number_format, c.border = "0%", BORDER
    c_last = 13 + len(cats)

    # Budget vs actual by trip
    for c, h in zip(range(5, 10), ["Trip", "Destination", "Budget (SAR)", "Actual (SAR)", "% used"]):
        cell = summ.cell(13, c, h)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, t in enumerate(data.trips, 14):
        src = i - 9  # Trips sheet row
        vals = [f"=Trips!A{src}", f"=Trips!D{src}", f"=Trips!H{src}", f"=Trips!I{src}", f"=Trips!K{src}"]
        for c, v in zip(range(5, 10), vals):
            cell = summ.cell(i, c, v)
            cell.border = BORDER
        for c in (7, 8):
            summ.cell(i, c).number_format = "#,##0"
        summ.cell(i, 9).number_format = "0%"
    summ.conditional_formatting.add(f"I14:I{13 + len(data.trips)}",
                                    FormulaRule(formula=["$I14>1"], fill=red, font=Font(bold=True, color="9B1C1C")))

    bar = BarChart()
    bar.type, bar.grouping, bar.title = "col", "clustered", "Budget vs actual by trip (SAR)"
    bar.add_data(Reference(summ, min_col=7, max_col=8, min_row=13, max_row=13 + len(data.trips)), titles_from_data=True)
    bar.set_categories(Reference(summ, min_col=6, min_row=14, max_row=13 + len(data.trips)))
    bar.y_axis.numFmt = "#,##0"
    bar.y_axis.majorGridlines = None
    bar.y_axis.scaling.min = 0
    bar.height, bar.width = 7.5, 13
    bar.series[0].graphicalProperties.solidFill = "C9D3DF"
    bar.series[1].graphicalProperties.solidFill = NAVY
    bar.y_axis.delete = False
    bar.x_axis.delete = False
    summ.add_chart(bar, f"E{c_last + 2}")

    cat_chart = BarChart()
    cat_chart.type, cat_chart.title = "bar", "Spend by category (SAR)"
    cat_chart.add_data(Reference(summ, min_col=2, min_row=13, max_row=c_last), titles_from_data=True)
    cat_chart.set_categories(Reference(summ, min_col=1, min_row=14, max_row=c_last))
    cat_chart.legend = None
    cat_chart.series[0].graphicalProperties.solidFill = TEAL
    cat_chart.x_axis.scaling.orientation = "maxMin"
    cat_chart.y_axis.numFmt = "#,##0"
    cat_chart.y_axis.majorGridlines = None
    cat_chart.y_axis.scaling.min = 0
    cat_chart.y_axis.delete = False
    cat_chart.x_axis.delete = False
    cat_chart.height, cat_chart.width = 7.5, 10.5
    summ.add_chart(cat_chart, f"A{c_last + 2}")

    for ws in (summ, exp, trips, pol, fx):
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    summ.page_setup.fitToHeight = 1
    summ.print_area = f"A1:I{c_last + 18}"
    out = Path(path)
    ensure_dir(out.parent)
    wb.save(out)
    return out


def generate(expenses_csv, trips_csv, rates_csv, policy_csv, out_dir, sample=True) -> dict:
    data = load(expenses_csv, trips_csv, rates_csv, policy_csv)
    out = ensure_dir(out_dir)
    xlsx = build_workbook(data, out / "expense_tracker.xlsx", sample)
    s = summarise(data)
    md = out / "expense_summary.md"
    rows = "\n".join(f"| {t['trip_id']} | {t['destination']} | {t['budget_sar']:,.0f} | {t['actual_sar']:,.0f} | "
                     f"{t['pct_used']:.0%} |" for t in s["trips"])
    flags = "\n".join(f"| {ln.date:%d %b} | {ln.trip_id} | {ln.description} | {ln.sar:,.2f} | {', '.join(ln.flags)} |"
                      for ln in s["flagged"])
    md.write_text(
        (f"> **{SAMPLE_BANNER}**\n\n" if sample else "")
        + "# Travel and expense summary\n\n" + text_summary(s) + "\n\n"
        + "## Budget vs actual\n\n| Trip | Destination | Budget (SAR) | Actual (SAR) | % used |\n|---|---|---|---|---|\n"
        + rows + "\n\n## Lines needing review\n\n| Date | Trip | Description | SAR | Flag |\n|---|---|---|---|---|\n"
        + flags + "\n", encoding="utf-8")
    return {"workbook": xlsx, "summary": md, "data": s}
